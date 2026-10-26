import io
import logging
import os
import re
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Tuple

import pymupdf as fitz  # PyMuPDF (modern import)
from docx import Document
from app.models.schemas import DocumentChunk

logger = logging.getLogger(__name__)

# Magic byte signatures
PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"

# Heading patterns for commercial agreements & contracts
HEADING_PATTERNS = [
    re.compile(r"^(?:ARTICLE|SECTION|CLAUSE|SCHEDULE|EXHIBIT)\s+([0-9IVXLCDM\.]+)[^\n]*", re.IGNORECASE),
    re.compile(r"^([0-9]{1,2}\.[0-9]{1,2}(?:\.[0-9]{1,2})?)\s+([A-Z][^\n]+)"),
    re.compile(r"^[A-Z0-9\s,\-\:]{4,60}$"),  # All-caps standalone headings
]


# Custom Ingestion Exceptions
class DocumentIngestionError(Exception):
    """Base exception for document ingestion failures."""
    pass


class FileValidationError(DocumentIngestionError):
    """Raised when file extension or magic bytes are invalid."""
    pass


class FileTooLargeError(DocumentIngestionError):
    """Raised when file size exceeds configured upload limit."""
    pass


class CorruptedFileError(DocumentIngestionError):
    """Raised when a file cannot be parsed due to corruption or structural damage."""
    pass


class ScannedPDFError(DocumentIngestionError):
    """Raised when a PDF contains scanned images but no extractable text layer."""
    pass


class EmptyDocumentError(DocumentIngestionError):
    """Raised when a document contains zero printable text content."""
    pass


class PathTraversalError(DocumentIngestionError):
    """Raised when a filename attempts directory traversal outside the uploads directory."""
    pass


def sanitize_filename(filename: str) -> str:
    """
    Sanitizes a filename to defend against path traversal, null bytes,
    and shell/filesystem injection. Returns only a safe base filename.
    """
    if not filename:
        return "unnamed_document"

    # Remove null bytes
    cleaned = filename.replace("\0", "")
    # Normalize forward/backward slashes to extract base name
    cleaned = cleaned.replace("\\", "/")
    cleaned = os.path.basename(cleaned)

    # Strip leading dots or traversal fragments
    cleaned = re.sub(r"^\.+", "", cleaned).strip()
    # Replace unsafe characters with underscore
    cleaned = re.sub(r"[^A-Za-z0-9._\-]", "_", cleaned)
    # Collapse multiple underscores
    cleaned = re.sub(r"_+", "_", cleaned)

    if not cleaned or cleaned in [".", ".."]:
        cleaned = "unnamed_document"

    # Limit maximum length
    if len(cleaned) > 120:
        stem, ext = os.path.splitext(cleaned)
        cleaned = stem[:120 - len(ext)] + ext

    return cleaned


def resolve_safe_path(base_dir: Path, filename: str) -> Path:
    """
    Resolves canonical path ensuring it resides strictly inside base_dir.
    Raises PathTraversalError if directory traversal is attempted.
    """
    base_dir_resolved = base_dir.resolve()
    target_path = (base_dir / filename).resolve()

    try:
        if not target_path.is_relative_to(base_dir_resolved):
            raise PathTraversalError(f"Path traversal detected: '{filename}' attempts to escape upload directory.")
    except AttributeError:
        if not str(target_path).startswith(str(base_dir_resolved)):
            raise PathTraversalError(f"Path traversal detected: '{filename}' attempts to escape upload directory.")

    return target_path


@dataclass
class DocumentSectionBlock:
    """Internal text block tracking page number, section header, and character bounds."""
    text: str
    page_number: int
    section: Optional[str] = None


class DocumentProcessor:
    """
    Enterprise document ingestion processor using PyMuPDF and python-docx.
    Features:
    - Strict validation (magic bytes, size, safe filenames, path traversal protection).
    - Preserves page numbers and section headings.
    - Scanned PDF detection (OCR warning).
    - Corrupted document and empty file handling with clear error messages.
    - Section-aware overlapping chunking with rich metadata.
    """

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 150):
        self.chunk_size = chunk_size
        self.chunk_overlap = min(chunk_overlap, chunk_size // 2)

    @staticmethod
    def validate_file_content(file_bytes: bytes, filename: str):
        """Validates file magic bytes against filename extension."""
        suffix = Path(filename).suffix.lower()

        if len(file_bytes) == 0:
            raise EmptyDocumentError("The uploaded file is empty (0 bytes).")

        if suffix == ".pdf":
            if not file_bytes.startswith(PDF_MAGIC):
                raise FileValidationError(
                    "Invalid PDF format: File header does not match standard PDF magic bytes (%PDF-)."
                )
        elif suffix in [".docx", ".doc"]:
            if not file_bytes.startswith(ZIP_MAGIC):
                raise FileValidationError(
                    "Invalid DOCX format: File header does not match OpenXML ZIP container signature."
                )
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    if not any(name.endswith(".xml") for name in zf.namelist()):
                        raise FileValidationError("Invalid DOCX format: Missing required OpenXML package components.")
            except zipfile.BadZipFile:
                raise CorruptedFileError("Corrupted DOCX file: Archive is incomplete or damaged.")
        elif suffix in [".txt", ".md"]:
            try:
                file_bytes[:1024].decode("utf-8")
            except UnicodeDecodeError:
                try:
                    file_bytes[:1024].decode("latin-1")
                except Exception:
                    raise FileValidationError("Malformed text document: Unable to decode characters.")
        else:
            raise FileValidationError(
                f"Unsupported document format '{suffix}'. Supported formats: .pdf, .docx, .txt"
            )

    def extract_structured_blocks(self, file_path: Path, filename: str) -> Tuple[List[DocumentSectionBlock], str]:
        """
        Extracts structured blocks preserving page numbers and section headings.
        Returns:
            (blocks: List[DocumentSectionBlock], full_text: str)
        """
        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            return self._extract_pdf_blocks(file_path)
        elif suffix in [".docx", ".doc"]:
            return self._extract_docx_blocks(file_path)
        elif suffix in [".txt", ".md"]:
            return self._extract_txt_blocks(file_path)
        else:
            raise FileValidationError(f"Unsupported document format '{suffix}'.")

    def _extract_pdf_blocks(self, file_path: Path) -> Tuple[List[DocumentSectionBlock], str]:
        """Extracts text page by page using PyMuPDF (fitz)."""
        blocks: List[DocumentSectionBlock] = []
        full_text_parts: List[str] = []

        try:
            doc = fitz.open(str(file_path))
        except Exception as e:
            raise CorruptedFileError(f"Corrupted PDF file: Unable to open or parse document ({str(e)}).")

        try:
            if doc.page_count == 0:
                raise EmptyDocumentError("The PDF document contains 0 pages.")

            total_extracted_chars = 0
            total_images_found = 0
            current_section = "Preamble"

            for page_num in range(1, doc.page_count + 1):
                try:
                    page = doc.load_page(page_num - 1)
                    text = page.get_text("text") or ""
                    images = page.get_images()
                    total_images_found += len(images)
                except Exception as e:
                    raise CorruptedFileError(f"Corrupted PDF page {page_num}: Error extracting content ({str(e)}).")

                cleaned_page_text = text.strip()
                total_extracted_chars += len(cleaned_page_text)

                if cleaned_page_text:
                    lines = cleaned_page_text.splitlines()
                    page_paragraphs: List[str] = []

                    for line in lines:
                        stripped = line.strip()
                        if not stripped:
                            continue

                        for pattern in HEADING_PATTERNS:
                            if pattern.match(stripped):
                                current_section = stripped[:80]
                                break

                        page_paragraphs.append(stripped)

                    page_text_content = "\n".join(page_paragraphs)
                    blocks.append(
                        DocumentSectionBlock(
                            text=page_text_content,
                            page_number=page_num,
                            section=current_section,
                        )
                    )
                    full_text_parts.append(f"--- [Page {page_num}] ---\n{page_text_content}")

            # Scanned PDF Detection: contains images but no extractable text
            if total_extracted_chars < 30 and total_images_found > 0:
                raise ScannedPDFError(
                    "Scanned PDF detected: The document contains scanned pages/images with no extractable text layer. "
                    "Please run OCR or provide a text-searchable PDF."
                )

            if not blocks or total_extracted_chars == 0:
                raise EmptyDocumentError("The PDF document contains no printable text.")

            return blocks, "\n\n".join(full_text_parts)
        finally:
            doc.close()

    def _extract_docx_blocks(self, file_path: Path) -> Tuple[List[DocumentSectionBlock], str]:
        """Extracts text paragraphs and headings using python-docx."""
        try:
            doc = Document(str(file_path))
        except Exception as e:
            raise CorruptedFileError(f"Corrupted DOCX file: Failed to load package ({str(e)}).")

        blocks: List[DocumentSectionBlock] = []
        full_text_parts: List[str] = []
        current_section = "General Provisions"
        page_estimate = 1
        chars_on_page = 0
        CHARS_PER_PAGE_ESTIMATE = 2500
        total_text_length = 0

        for para in doc.paragraphs:
            raw_text = para.text.strip()
            if not raw_text:
                continue

            total_text_length += len(raw_text)

            style_name = para.style.name.lower() if para.style else ""
            if "heading" in style_name or "title" in style_name:
                current_section = raw_text[:80]
            else:
                for pattern in HEADING_PATTERNS:
                    if pattern.match(raw_text):
                        current_section = raw_text[:80]
                        break

            chars_on_page += len(raw_text)
            if chars_on_page > CHARS_PER_PAGE_ESTIMATE:
                page_estimate += 1
                chars_on_page = 0

            blocks.append(
                DocumentSectionBlock(
                    text=raw_text,
                    page_number=page_estimate,
                    section=current_section,
                )
            )
            full_text_parts.append(raw_text)

        # Extract table data
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_rows.append(row_text)
            if table_rows:
                table_content = "\n".join(table_rows)
                total_text_length += len(table_content)
                blocks.append(
                    DocumentSectionBlock(
                        text=table_content,
                        page_number=page_estimate,
                        section=f"{current_section} (Table Data)",
                    )
                )
                full_text_parts.append(table_content)

        if total_text_length == 0 or not blocks:
            raise EmptyDocumentError("The DOCX document contains no text content.")

        return blocks, "\n\n".join(full_text_parts)

    def _extract_txt_blocks(self, file_path: Path) -> Tuple[List[DocumentSectionBlock], str]:
        """Extracts text from plain text or markdown files."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        except Exception as e:
            raise CorruptedFileError(f"Error reading text document: {str(e)}")

        stripped = content.strip()
        if not stripped:
            raise EmptyDocumentError("The text document contains no content.")

        blocks: List[DocumentSectionBlock] = []
        paragraphs = stripped.split("\n\n")
        current_section = "Agreement Text"
        page_estimate = 1
        chars_on_page = 0

        for p in paragraphs:
            text = p.strip()
            if not text:
                continue

            for pattern in HEADING_PATTERNS:
                if pattern.match(text):
                    current_section = text[:80]
                    break

            chars_on_page += len(text)
            if chars_on_page > 2500:
                page_estimate += 1
                chars_on_page = 0

            blocks.append(
                DocumentSectionBlock(
                    text=text,
                    page_number=page_estimate,
                    section=current_section,
                )
            )

        return blocks, content

    def chunk_document(
        self,
        document_id: str,
        filename: str,
        blocks: List[DocumentSectionBlock],
        full_text: str,
    ) -> List[DocumentChunk]:
        """
        Splits extracted structured blocks into section-aware overlapping chunks.
        Guarantees loop termination, forward progress, and metadata retention.
        """
        chunks: List[DocumentChunk] = []
        if not blocks:
            return chunks

        chunk_index = 0
        current_buffer: List[str] = []
        current_length = 0
        current_page = blocks[0].page_number
        current_section = blocks[0].section or "General"
        running_char_offset = 0

        for block in blocks:
            # If section changed and buffer already has substantial text, flush chunk
            if block.section and block.section != current_section and current_length >= (self.chunk_size // 2):
                chunk_text = "\n\n".join(current_buffer).strip()
                if chunk_text:
                    chunk = self._create_chunk(
                        document_id=document_id,
                        filename=filename,
                        chunk_index=chunk_index,
                        text=chunk_text,
                        start_char=running_char_offset,
                        page_number=current_page,
                        section=current_section,
                    )
                    chunks.append(chunk)
                    chunk_index += 1
                    running_char_offset += len(chunk_text)

                overlap_text = chunk_text[-self.chunk_overlap:] if len(chunk_text) > self.chunk_overlap else ""
                current_buffer = [overlap_text, block.text] if overlap_text else [block.text]
                current_length = len(overlap_text) + len(block.text)
                current_page = block.page_number
                current_section = block.section or current_section
                continue

            current_buffer.append(block.text)
            current_length += len(block.text)
            current_page = block.page_number
            if block.section:
                current_section = block.section

            # When buffer exceeds chunk size, build chunk and slide window with overlap
            while current_length >= self.chunk_size:
                combined = "\n\n".join(current_buffer).strip()
                split_point = self._find_optimal_split(combined, self.chunk_size)

                # Ensure forward progress
                min_advance = max(self.chunk_overlap + 10, self.chunk_size // 2)
                if split_point < min_advance:
                    split_point = min(self.chunk_size, len(combined))

                chunk_text = combined[:split_point].strip()
                if chunk_text:
                    chunk = self._create_chunk(
                        document_id=document_id,
                        filename=filename,
                        chunk_index=chunk_index,
                        text=chunk_text,
                        start_char=running_char_offset,
                        page_number=current_page,
                        section=current_section,
                    )
                    chunks.append(chunk)
                    chunk_index += 1
                    running_char_offset += len(chunk_text)

                # Slide window
                overlap_start = max(0, split_point - self.chunk_overlap)
                remaining = combined[overlap_start:].strip()

                if len(remaining) >= len(combined):
                    remaining = combined[split_point:].strip()

                current_buffer = [remaining] if remaining else []
                current_length = len(remaining)

        # Flush any remaining buffer text
        if current_buffer:
            trailing_text = "\n\n".join(current_buffer).strip()
            if trailing_text:
                chunk = self._create_chunk(
                    document_id=document_id,
                    filename=filename,
                    chunk_index=chunk_index,
                    text=trailing_text,
                    start_char=running_char_offset,
                    page_number=current_page,
                    section=current_section,
                )
                chunks.append(chunk)

        return chunks

    def _create_chunk(
        self,
        document_id: str,
        filename: str,
        chunk_index: int,
        text: str,
        start_char: int,
        page_number: int,
        section: Optional[str],
    ) -> DocumentChunk:
        return DocumentChunk(
            chunk_id=str(uuid.uuid4()),
            document_id=document_id,
            contract_id=document_id,
            filename=filename,
            chunk_index=chunk_index,
            text=text,
            start_char=start_char,
            end_char=start_char + len(text),
            page_number=page_number,
            section=section,
            section_title=section,
            char_count=len(text),
            word_count=len(text.split()),
        )

    @staticmethod
    def _find_optimal_split(text: str, target_size: int) -> int:
        """Finds natural paragraph, newline, or sentence split boundary near target_size."""
        if len(text) <= target_size:
            return len(text)

        search_window = text[max(0, target_size - 150):min(len(text), target_size + 100)]
        offset = max(0, target_size - 150)

        # Try paragraph break
        para_idx = search_window.rfind("\n\n")
        if para_idx != -1 and (offset + para_idx) > 20:
            return offset + para_idx

        # Try newline
        newline_idx = search_window.rfind("\n")
        if newline_idx != -1 and (offset + newline_idx) > 20:
            return offset + newline_idx

        # Try sentence end
        sentence_idx = search_window.rfind(". ")
        if sentence_idx != -1 and (offset + sentence_idx) > 20:
            return offset + sentence_idx + 1

        # Fallback to whitespace
        space_idx = search_window.rfind(" ")
        if space_idx != -1 and (offset + space_idx) > 20:
            return offset + space_idx

        return target_size


# Default singleton instance
document_processor = DocumentProcessor()
