import io
import os
import tempfile
from pathlib import Path
import pytest
import fitz  # PyMuPDF
from docx import Document

from app.services.document_processor import (
    DocumentProcessor,
    DocumentSectionBlock,
    CorruptedFileError,
    EmptyDocumentError,
    FileValidationError,
    PathTraversalError,
    ScannedPDFError,
    resolve_safe_path,
    sanitize_filename,
)


class TestFilenameAndSecurity:
    """Tests for filename sanitization and path traversal defenses."""

    def test_sanitize_filename_normal(self):
        assert sanitize_filename("Master_Agreement.pdf") == "Master_Agreement.pdf"

    def test_sanitize_filename_traversal_attempts(self):
        assert sanitize_filename("../../etc/passwd.pdf") == "passwd.pdf"
        assert sanitize_filename("..\\..\\Windows\\System32\\calc.docx") == "calc.docx"
        assert sanitize_filename("../../../secret.pdf") == "secret.pdf"

    def test_sanitize_filename_null_bytes(self):
        sanitized = sanitize_filename("contract\0evil.pdf")
        assert "\0" not in sanitized
        assert "contractevil.pdf" in sanitized

    def test_sanitize_filename_special_characters(self):
        sanitized = sanitize_filename("Contract @#$$% 2026! (Draft).pdf")
        assert "@" not in sanitized
        assert "$" not in sanitized
        assert "Contract_2026_Draft_.pdf" in sanitized

    def test_resolve_safe_path_valid(self, tmp_path):
        safe_path = resolve_safe_path(tmp_path, "contract_123.pdf")
        assert safe_path.parent == tmp_path.resolve()

    def test_resolve_safe_path_traversal_rejected(self, tmp_path):
        with pytest.raises(PathTraversalError):
            resolve_safe_path(tmp_path, "../outside.txt")


class TestMagicBytesValidation:
    """Tests for magic byte content inspection."""

    def test_valid_pdf_magic(self):
        valid_pdf_bytes = b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n"
        DocumentProcessor.validate_file_content(valid_pdf_bytes, "test.pdf")

    def test_invalid_pdf_magic_rejected(self):
        invalid_bytes = b"NOT_A_PDF_FILE"
        with pytest.raises(FileValidationError) as exc:
            DocumentProcessor.validate_file_content(invalid_bytes, "fake.pdf")
        assert "magic bytes" in str(exc.value)

    def test_empty_bytes_rejected(self):
        with pytest.raises(EmptyDocumentError):
            DocumentProcessor.validate_file_content(b"", "empty.pdf")

    def test_invalid_docx_magic_rejected(self):
        invalid_bytes = b"PLAIN_TEXT_NOT_ZIP"
        with pytest.raises(FileValidationError):
            DocumentProcessor.validate_file_content(invalid_bytes, "fake.docx")


class TestPdfParsingAndMetadata:
    """Tests for PDF text extraction and metadata preservation using PyMuPDF."""

    @pytest.fixture
    def sample_pdf_path(self, tmp_path):
        """Creates a real 2-page PDF document with sections."""
        pdf_file = tmp_path / "commercial_agreement.pdf"
        doc = fitz.open()

        # Page 1
        page1 = doc.new_page()
        page1.insert_text(
            (50, 50),
            "ARTICLE 1. INDEMNITY\nProvider agrees to indemnify customer against all third-party claims.",
            fontsize=12,
        )

        # Page 2
        page2 = doc.new_page()
        page2.insert_text(
            (50, 50),
            "ARTICLE 2. TERMINATION\nEither party may terminate upon thirty days written notice.",
            fontsize=12,
        )

        doc.save(str(pdf_file))
        doc.close()
        return pdf_file

    def test_pdf_extraction_preserves_pages_and_sections(self, sample_pdf_path):
        processor = DocumentProcessor(chunk_size=500, chunk_overlap=50)
        blocks, full_text = processor.extract_structured_blocks(sample_pdf_path, sample_pdf_path.name)

        assert len(blocks) == 2
        assert blocks[0].page_number == 1
        assert "ARTICLE 1" in blocks[0].section
        assert "indemnify" in blocks[0].text

        assert blocks[1].page_number == 2
        assert "ARTICLE 2" in blocks[1].section
        assert "terminate" in blocks[1].text

        assert "--- [Page 1] ---" in full_text
        assert "--- [Page 2] ---" in full_text

    def test_pdf_chunking_metadata(self, sample_pdf_path):
        processor = DocumentProcessor(chunk_size=80, chunk_overlap=15)
        blocks, full_text = processor.extract_structured_blocks(sample_pdf_path, sample_pdf_path.name)
        chunks = processor.chunk_document(
            document_id="doc-test-1",
            filename=sample_pdf_path.name,
            blocks=blocks,
            full_text=full_text,
        )

        assert len(chunks) >= 2
        chunk1 = chunks[0]
        assert chunk1.document_id == "doc-test-1"
        assert chunk1.contract_id == "doc-test-1"
        assert chunk1.filename == sample_pdf_path.name
        assert chunk1.page_number in [1, 2]
        assert chunk1.section is not None
        assert chunk1.char_count > 0
        assert chunk1.word_count > 0
        assert chunk1.start_char >= 0
        assert chunk1.end_char > chunk1.start_char

    def test_empty_pdf_detection(self, tmp_path):
        pdf_file = tmp_path / "empty.pdf"
        doc = fitz.open()
        doc.new_page()  # Page with no text
        doc.save(str(pdf_file))
        doc.close()

        processor = DocumentProcessor()
        with pytest.raises(EmptyDocumentError):
            processor.extract_structured_blocks(pdf_file, pdf_file.name)

    def test_scanned_pdf_detection(self, tmp_path):
        """Simulate a scanned PDF: has an image but 0 extractable text."""
        pdf_file = tmp_path / "scanned_invoice.pdf"
        doc = fitz.open()
        page = doc.new_page()

        # Insert a 100x100 white pixel pixmap
        pix = fitz.Pixmap(fitz.csRGB, (0, 0, 100, 100), 1)
        pix.clear_with(255)
        img_bytes = pix.tobytes("png")
        page.insert_image(page.rect, stream=img_bytes)

        doc.save(str(pdf_file))
        doc.close()

        processor = DocumentProcessor()
        with pytest.raises(ScannedPDFError) as exc:
            processor.extract_structured_blocks(pdf_file, pdf_file.name)
        assert "Scanned PDF detected" in str(exc.value)

    def test_corrupted_pdf_error(self, tmp_path):
        corrupted_file = tmp_path / "broken.pdf"
        with open(corrupted_file, "wb") as f:
            f.write(b"%PDF-1.7\nCorrupted binary payload\x00\xff truncated")

        processor = DocumentProcessor()
        with pytest.raises(CorruptedFileError) as exc:
            processor.extract_structured_blocks(corrupted_file, corrupted_file.name)
        assert "Corrupted PDF" in str(exc.value)


class TestDocxParsingAndMetadata:
    """Tests for Word DOCX extraction and heading detection using python-docx."""

    @pytest.fixture
    def sample_docx_path(self, tmp_path):
        docx_file = tmp_path / "sample_contract.docx"
        doc = Document()
        doc.add_heading("ARTICLE 1. SCOPE OF SERVICES", level=1)
        doc.add_paragraph("Provider will perform software engineering and cloud deployment services.")
        doc.add_heading("ARTICLE 2. LIMITATION OF LIABILITY", level=1)
        doc.add_paragraph("Customer aggregate liability shall not exceed total fees paid in the prior six months.")

        doc.save(str(docx_file))
        return docx_file

    def test_docx_extraction_and_sections(self, sample_docx_path):
        processor = DocumentProcessor()
        blocks, full_text = processor.extract_structured_blocks(sample_docx_path, sample_docx_path.name)

        assert len(blocks) >= 2
        sections = [b.section for b in blocks if b.section]
        assert any("ARTICLE 1" in s for s in sections)
        assert any("ARTICLE 2" in s for s in sections)
        assert "cloud deployment" in full_text

    def test_corrupted_docx_error(self, tmp_path):
        corrupted = tmp_path / "broken.docx"
        with open(corrupted, "wb") as f:
            f.write(b"PK\x03\x04corrupted zip payload not ooxml")

        processor = DocumentProcessor()
        with pytest.raises(CorruptedFileError):
            processor.extract_structured_blocks(corrupted, corrupted.name)


class TestChunkingAlgorithm:
    """Tests for section-aware overlapping chunk logic."""

    def test_overlapping_chunks_preserve_context(self):
        processor = DocumentProcessor(chunk_size=100, chunk_overlap=30)
        blocks = [
            DocumentSectionBlock(
                text="The quick brown fox jumps over the lazy dog repeatedly to generate substantial characters.",
                page_number=1,
                section="SECTION 1. FOX TEST",
            ),
            DocumentSectionBlock(
                text="Second paragraph continuing the narrative of the agile animal traversing the domestic pet.",
                page_number=1,
                section="SECTION 1. FOX TEST",
            ),
            DocumentSectionBlock(
                text="SECTION 2. NEW CLAUSE\nThis starts an entirely fresh section with its own distinct heading.",
                page_number=2,
                section="SECTION 2. NEW CLAUSE",
            ),
        ]

        chunks = processor.chunk_document(
            document_id="doc-99",
            filename="fox.txt",
            blocks=blocks,
            full_text="...",
        )

        assert len(chunks) >= 2
        # Check that chunk indices are sequential
        for i, chunk in enumerate(chunks):
            assert chunk.chunk_index == i
            assert chunk.document_id == "doc-99"
            assert chunk.filename == "fox.txt"
            assert chunk.char_count > 0
            assert chunk.word_count > 0
